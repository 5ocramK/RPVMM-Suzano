import openpyxl
import os
import shutil
import csv

BASE_DIR = r"c:\Users\Marco\Desktop\Atividades\RPVMM"
EXCEL_PATH = os.path.join(BASE_DIR, "dados_limpos.xlsm")
EXCEL_ORIG_PATH = os.path.join(BASE_DIR, "dados.xlsx")
SUZANO_DIR = os.path.join(BASE_DIR, "RPVMM-Suzano")
BACKUP_DIR = os.path.join(SUZANO_DIR, "backup_csv")

os.makedirs(BACKUP_DIR, exist_ok=True)

# 1. Backup existing CSVs if not already backed up
csv_files = [f for f in os.listdir(SUZANO_DIR) if f.endswith('.csv')]
print(f"Checking backup for {len(csv_files)} existing CSVs in {BACKUP_DIR}...")
for f in csv_files:
    dst = os.path.join(BACKUP_DIR, f)
    if not os.path.exists(dst):
        shutil.copy2(os.path.join(SUZANO_DIR, f), dst)

# 2. Sheet mappings and column specifications
# Note: MDLocais is extracted from master list (dados.xlsx) to ensure all PODs
# referenced by multi-stop routes and initial inventory are present.
mappings = {
    'MDClientes.csv': ('MDClientes', ['COD_CLIENTE', 'ATIVO'], EXCEL_PATH),
    'MDArmadorRotas.csv': ('MDArmadorRotas - LIMPO', ['COD_ARMADOR', 'COD_LOCAL_ORIGENS', 'COD_LOCAL_DESTINOS', 'MIN_INTAKE', 'MAX_INTAKE', 'FRETE_POR_TON', 'TEMPO_IDA', 'TEMPO_IDA_VOLTA'], EXCEL_PATH),
    'MDLocais.csv': ('MDLocais', ['COD_LOCAL', 'MAX_CARREGAMENTO', 'MAX_DESCARREGAMENTO', 'MAX_ESTOQUE', 'CALADO_DWT', 'TIPO'], EXCEL_ORIG_PATH),
    'RDCustoInland.csv': ('RDCustoInland - LIMPO', ['COD_CLIENTE', 'COD_LOCAL', 'CUSTO_INLAND_POR_TON'], EXCEL_PATH),
    'RDDemanda.csv': ('RDDemanda', ['COD_CLIENTE', 'FAMILIA_PRODUTO', 'LINHA_PRODUTO', 'INSTANTE', 'PRECO_VENDA', 'DEMANDA'], EXCEL_PATH),
    'RDEstoqueInicial.csv': ('RDEstoqueInicial', ['COD_LOCAL', 'FAMILIA_PRODUTO', 'LINHA_PRODUTO', 'QTD_DISPONIVEL'], EXCEL_PATH),
    'MDProdutos.csv': ('MDProdutos', ['FAMILIA_PRODUTO', 'LINHA_PRODUTO', 'ATIVO'], EXCEL_PATH),
    'RDCapFabricas.csv': ('RDCapFabricas', ['COD_LOCAL', 'INSTANTE', 'VOLUME_PRODUCAO'], EXCEL_PATH),
    'RDLocalProduto.csv': ('RDLocalProduto', ['COD_LOCAL', 'LINHA_PRODUTO', 'PRODUZ'], EXCEL_PATH)
}

def format_val(v):
    if v is None:
        return ''
    if isinstance(v, bool):
        return 'TRUE' if v else 'FALSE'
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        # Format float with comma decimal representation matching original CSV style
        return str(v).replace('.', ',')
    if isinstance(v, int):
        return str(v)
    return str(v).strip()

print(f"Loading workbooks...")
wb_limpos = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
wb_orig = openpyxl.load_workbook(EXCEL_ORIG_PATH, read_only=True, data_only=True) if os.path.exists(EXCEL_ORIG_PATH) else None

for csv_name, (sheet_name, target_cols, file_src) in mappings.items():
    wb = wb_orig if file_src == EXCEL_ORIG_PATH and wb_orig is not None else wb_limpos
    ws = wb[sheet_name]
    header = None
    col_indices = []
    rows = []
    for i, r in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            header = list(r)
            col_indices = [header.index(col) for col in target_cols]
            continue
        # Stop or skip if primary key/first column is null or empty
        if r[col_indices[0]] is None or str(r[col_indices[0]]).strip() == '':
            continue
        formatted_row = [format_val(r[idx]) for idx in col_indices]
        rows.append(formatted_row)
    
    out_path = os.path.join(SUZANO_DIR, csv_name)
    with open(out_path, 'w', newline='', encoding='utf-8') as fp:
        writer = csv.writer(fp)
        writer.writerow(target_cols)
        writer.writerows(rows)
    
    print(f"Successfully generated {csv_name} from sheet '{sheet_name}': {len(rows)} records written.")

print("All CSVs converted and updated successfully.")
